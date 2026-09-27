import csv
from datetime import timedelta
from io import BytesIO

from django.contrib import messages
from django.contrib.auth.mixins import LoginRequiredMixin, UserPassesTestMixin
from django.db.models import Count, Q
from django.http import HttpResponse
from django.utils import timezone
from django.urls import reverse_lazy
from django.views import View
from django.views.generic import CreateView, DeleteView, DetailView, ListView, TemplateView, UpdateView
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, letter
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle

from .forms import ContactMessageForm, DriverForm, VehicleForm
from .models import Driver, Vehicle


EXPORT_HEADERS = [
    'Vehicle Number',
    'Make',
    'Model',
    'Year',
    'Type',
    'Owner',
    'Driver',
    'Status',
    'Last Service',
    'Next Service',
    'Service Cost',
    'Repair Notes',
]


def get_export_rows():
    rows = []
    vehicles = Vehicle.objects.select_related('driver').order_by('vehicle_number')
    for vehicle in vehicles:
        rows.append([
            vehicle.vehicle_number,
            vehicle.make,
            vehicle.model,
            vehicle.year,
            vehicle.vehicle_type,
            vehicle.owner_name,
            vehicle.driver.name if vehicle.driver else 'Not assigned',
            vehicle.get_status_display(),
            vehicle.last_service_date.strftime('%Y-%m-%d') if vehicle.last_service_date else '',
            vehicle.next_service_date.strftime('%Y-%m-%d') if vehicle.next_service_date else '',
            str(vehicle.service_cost or ''),
            vehicle.repair_notes,
        ])
    return rows


class StaffRequiredMixin(LoginRequiredMixin, UserPassesTestMixin):
    login_url = 'login'

    def test_func(self):
        user = self.request.user
        return user.is_superuser or user.groups.filter(name__in=['Admin', 'Staff']).exists()


class AdminRequiredMixin(LoginRequiredMixin, UserPassesTestMixin):
    login_url = 'login'

    def test_func(self):
        user = self.request.user
        return user.is_superuser or user.groups.filter(name='Admin').exists()


class DashboardView(TemplateView):
    template_name = 'vehicles/dashboard.html'
    
    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        total_count = Vehicle.objects.count()
        available_count = Vehicle.objects.filter(status=Vehicle.STATUS_AVAILABLE).count()
        assigned_count = Vehicle.objects.filter(status=Vehicle.STATUS_ASSIGNED).count()
        maintenance_count = Vehicle.objects.filter(status=Vehicle.STATUS_MAINTENANCE).count()
        type_counts = list(
            Vehicle.objects.values('vehicle_type')
            .annotate(count=Count('id'))
            .order_by('-count', 'vehicle_type')[:6]
        )
        max_type_count = max([item['count'] for item in type_counts], default=1)

        context['total_count'] = total_count
        context['available_count'] = available_count
        context['assigned_count'] = assigned_count
        context['maintenance_count'] = maintenance_count
        context['status_chart'] = [
            {
                'label': 'Available',
                'count': available_count,
                'percent': round((available_count / total_count) * 100) if total_count else 0,
                'class': 'available',
            },
            {
                'label': 'Assigned',
                'count': assigned_count,
                'percent': round((assigned_count / total_count) * 100) if total_count else 0,
                'class': 'assigned',
            },
            {
                'label': 'Maintenance',
                'count': maintenance_count,
                'percent': round((maintenance_count / total_count) * 100) if total_count else 0,
                'class': 'maintenance',
            },
        ]
        context['type_chart'] = [
            {
                'label': item['vehicle_type'],
                'count': item['count'],
                'percent': round((item['count'] / max_type_count) * 100),
            }
            for item in type_counts
        ]
        context['recent_vehicles'] = Vehicle.objects.select_related('driver').order_by('-created_at')[:5]
        return context


class VehicleListView(ListView):
    model = Vehicle
    context_object_name = 'vehicles'
    template_name = 'vehicles/vehicle_list.html'

    def get_queryset(self):
        queryset = super().get_queryset().select_related('driver')
        query = self.request.GET.get('q', '').strip()
        status = self.request.GET.get('status', '').strip()
        if query:
            queryset = queryset.filter(
                Q(vehicle_number__icontains=query)
                | Q(make__icontains=query)
                | Q(model__icontains=query)
                | Q(owner_name__icontains=query)
                | Q(driver__name__icontains=query)
            )
        if status in dict(Vehicle.STATUS_CHOICES):
            queryset = queryset.filter(status=status)
        return queryset

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['query'] = self.request.GET.get('q', '').strip()
        context['selected_status'] = self.request.GET.get('status', '').strip()
        context['status_choices'] = Vehicle.STATUS_CHOICES
        return context


class ContactView(TemplateView):
    template_name = 'vehicles/contact.html'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['form'] = kwargs.get('form') or ContactMessageForm()
        return context

    def post(self, request, *args, **kwargs):
        form = ContactMessageForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, 'Thanks, your message has been saved.')
            return self.get(request, *args, **kwargs)
        return self.render_to_response(self.get_context_data(form=form))


class DriverListView(ListView):
    model = Driver
    context_object_name = 'drivers'
    template_name = 'vehicles/driver_list.html'

    def get_queryset(self):
        queryset = Driver.objects.prefetch_related('vehicles').order_by('name')
        query = self.request.GET.get('q', '').strip()
        availability = self.request.GET.get('availability', '').strip()
        if query:
            queryset = queryset.filter(
                Q(name__icontains=query)
                | Q(license_number__icontains=query)
                | Q(phone__icontains=query)
            )
        if availability == 'available':
            queryset = queryset.filter(is_available=True)
        elif availability == 'unavailable':
            queryset = queryset.filter(is_available=False)
        return queryset

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['query'] = self.request.GET.get('q', '').strip()
        context['selected_availability'] = self.request.GET.get('availability', '').strip()
        context['total_driver_count'] = Driver.objects.count()
        context['available_driver_count'] = Driver.objects.filter(is_available=True).count()
        context['unavailable_driver_count'] = Driver.objects.filter(is_available=False).count()
        return context


class DriverCreateView(StaffRequiredMixin, CreateView):
    model = Driver
    form_class = DriverForm
    template_name = 'vehicles/driver_form.html'
    success_url = reverse_lazy('driver-list')

    def form_valid(self, form):
        messages.success(self.request, 'Driver added successfully.')
        return super().form_valid(form)


class DriverUpdateView(StaffRequiredMixin, UpdateView):
    model = Driver
    form_class = DriverForm
    template_name = 'vehicles/driver_form.html'
    success_url = reverse_lazy('driver-list')

    def form_valid(self, form):
        messages.success(self.request, 'Driver updated successfully.')
        return super().form_valid(form)


class DriverDeleteView(AdminRequiredMixin, DeleteView):
    model = Driver
    template_name = 'vehicles/driver_confirm_delete.html'
    success_url = reverse_lazy('driver-list')

    def form_valid(self, form):
        messages.success(self.request, 'Driver deleted successfully.')
        return super().form_valid(form)


class MaintenanceView(ListView):
    model = Vehicle
    context_object_name = 'maintenance_vehicles'
    template_name = 'vehicles/maintenance.html'

    def get_queryset(self):
        return (
            Vehicle.objects.select_related('driver')
            .filter(Q(status=Vehicle.STATUS_MAINTENANCE) | Q(next_service_date__isnull=False))
            .order_by('next_service_date', 'vehicle_number')
        )

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        today = timezone.localdate()
        due_limit = today + timedelta(days=30)
        service_costs = [
            vehicle.service_cost or 0
            for vehicle in context['maintenance_vehicles']
        ]
        context['today'] = today
        context['due_soon_count'] = Vehicle.objects.filter(
            next_service_date__gte=today,
            next_service_date__lte=due_limit,
        ).count()
        context['maintenance_count'] = Vehicle.objects.filter(status=Vehicle.STATUS_MAINTENANCE).count()
        context['total_service_cost'] = sum(service_costs)
        return context


class VehicleDetailView(DetailView):
    model = Vehicle
    context_object_name = 'vehicle'
    template_name = 'vehicles/vehicle_detail.html'


class VehicleCreateView(StaffRequiredMixin, CreateView):
    model = Vehicle
    form_class = VehicleForm
    template_name = 'vehicles/vehicle_form.html'
    success_url = reverse_lazy('vehicle-list')

    def form_valid(self, form):
        messages.success(self.request, 'Vehicle added successfully.')
        return super().form_valid(form)


class VehicleUpdateView(StaffRequiredMixin, UpdateView):
    model = Vehicle
    form_class = VehicleForm
    template_name = 'vehicles/vehicle_form.html'
    success_url = reverse_lazy('vehicle-list')

    def form_valid(self, form):
        messages.success(self.request, 'Vehicle updated successfully.')
        return super().form_valid(form)


class VehicleDeleteView(AdminRequiredMixin, DeleteView):
    model = Vehicle
    template_name = 'vehicles/vehicle_confirm_delete.html'
    success_url = reverse_lazy('vehicle-list')

    def form_valid(self, form):
        messages.success(self.request, 'Vehicle deleted successfully.')
        return super().form_valid(form)


class VehicleCSVExportView(StaffRequiredMixin, View):
    def get(self, request):
        response = HttpResponse(content_type='text/csv')
        response['Content-Disposition'] = 'attachment; filename=\"vehicle_records.csv\"'
        writer = csv.writer(response)
        writer.writerow(EXPORT_HEADERS)
        writer.writerows(get_export_rows())
        return response


class VehicleExcelExportView(StaffRequiredMixin, View):
    def get(self, request):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'Vehicle Records'
        sheet.append(EXPORT_HEADERS)

        header_fill = PatternFill('solid', fgColor='D9EAF7')
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill

        for row in get_export_rows():
            sheet.append(row)

        for column_cells in sheet.columns:
            width = max(len(str(cell.value or '')) for cell in column_cells)
            sheet.column_dimensions[column_cells[0].column_letter].width = min(width + 2, 42)

        stream = BytesIO()
        workbook.save(stream)
        stream.seek(0)

        response = HttpResponse(
            stream.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        response['Content-Disposition'] = 'attachment; filename=\"vehicle_records.xlsx\"'
        return response


class VehiclePDFExportView(StaffRequiredMixin, View):
    def get(self, request):
        stream = BytesIO()
        document = SimpleDocTemplate(stream, pagesize=landscape(letter))
        table_data = [EXPORT_HEADERS[:10]] + [row[:10] for row in get_export_rows()]
        table = Table(table_data, repeatRows=1)
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#d9eaf7')),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 8),
            ('GRID', (0, 0), (-1, -1), 0.25, colors.HexColor('#c8d1df')),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ]))
        document.build([table])
        stream.seek(0)

        response = HttpResponse(stream.getvalue(), content_type='application/pdf')
        response['Content-Disposition'] = 'attachment; filename=\"vehicle_records.pdf\"'
        return response
